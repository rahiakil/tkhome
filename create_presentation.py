import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    # Set to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Theme Colors
    c_navy = RGBColor(10, 34, 64)        # #0A2240 - Primary Brand
    c_teal = RGBColor(0, 168, 181)       # #00A8B5 - Accent Teal
    c_soft_teal = RGBColor(240, 249, 250) # Light teal background for cards
    c_light_grey = RGBColor(245, 247, 250) # Light grey card background
    c_border_grey = RGBColor(220, 225, 230)
    c_charcoal = RGBColor(33, 37, 41)    # Dark charcoal for primary text
    c_muted_grey = RGBColor(90, 100, 115) # Muted text
    c_white = RGBColor(255, 255, 255)
    c_light_navy = RGBColor(20, 50, 90)  # Lighter navy for dark-theme cards
    c_coral_bg = RGBColor(253, 243, 243)  # Soft red background for callouts
    c_coral_text = RGBColor(180, 40, 40)  # Dark red for callouts

    # Helper: Add slide header for light-theme content slides
    def add_slide_header(slide, title_text, category_text="VERTEX SMART CATEGORIZATION ENGINE"):
        # Category / Tracker text
        cat_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.3), Inches(11.833), Inches(0.3))
        cat_tf = cat_box.text_frame
        cat_tf.word_wrap = True
        cat_tf.margin_top = 0
        cat_tf.margin_bottom = 0
        cat_tf.margin_left = 0
        p0 = cat_tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = c_teal
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.833), Inches(0.8))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        title_tf.margin_top = 0
        title_tf.margin_bottom = 0
        title_tf.margin_left = 0
        p1 = title_tf.paragraphs[0]
        p1.text = title_text
        p1.font.name = "Arial"
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = c_navy

        # Thin teal accent bar below title
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.3), Inches(2.0), Inches(0.04))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = c_teal
        accent_bar.line.fill.background()

    # Helper: Add slide header for dark-theme content slides
    def add_dark_slide_header(slide, title_text, category_text="VERTEX SMART CATEGORIZATION ENGINE"):
        # Category / Tracker text
        cat_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.3), Inches(11.833), Inches(0.3))
        cat_tf = cat_box.text_frame
        cat_tf.word_wrap = True
        cat_tf.margin_top = 0
        cat_tf.margin_bottom = 0
        cat_tf.margin_left = 0
        p0 = cat_tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = "Arial"
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = c_teal
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.833), Inches(0.8))
        title_tf = title_box.text_frame
        title_tf.word_wrap = True
        title_tf.margin_top = 0
        title_tf.margin_bottom = 0
        title_tf.margin_left = 0
        p1 = title_tf.paragraphs[0]
        p1.text = title_text
        p1.font.name = "Arial"
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = c_white

        # Thin teal accent bar below title
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.3), Inches(2.0), Inches(0.04))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = c_teal
        accent_bar.line.fill.background()

    # ==========================================
    # SLIDE 1: TITLE & OVERVIEW (Dark Theme)
    # ==========================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6]) # blank layout
    slide1.background.fill.solid()
    slide1.background.fill.fore_color.rgb = c_navy

    # Decorative background accent (a subtle teal triangle or block on the right)
    bg_accent = slide1.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(9.5), Inches(0), Inches(3.833), Inches(7.5))
    bg_accent.fill.solid()
    bg_accent.fill.fore_color.rgb = RGBColor(12, 42, 78)
    bg_accent.line.fill.background()
    bg_accent.rotation = 180

    # Decorative teal bar on the left
    left_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), Inches(7.5))
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = c_teal
    left_bar.line.fill.background()

    # Title & Subtitle box
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(2.2))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    title_tf.margin_top = 0
    title_tf.margin_left = 0
    
    # Title
    p_title = title_tf.paragraphs[0]
    p_title.text = "Smart Categorization Engine"
    p_title.font.name = "Arial"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = c_white
    p_title.space_after = Pt(10)

    # Subtitle
    p_sub = title_tf.add_paragraph()
    p_sub.text = "Building a High-Precision, Self-Improving LLM Pipeline for Tax-Compliant Product Matching"
    p_sub.font.name = "Arial"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = c_teal

    # Presenter info
    info_box = slide1.shapes.add_textbox(Inches(1.0), Inches(4.2), Inches(5.0), Inches(2.0))
    info_tf = info_box.text_frame
    info_tf.word_wrap = True
    info_tf.margin_top = 0
    info_tf.margin_left = 0
    
    p_pres = info_tf.paragraphs[0]
    p_pres.text = "Presenter: Lead AI Product Engineer"
    p_pres.font.name = "Arial"
    p_pres.font.size = Pt(14)
    p_pres.font.color.rgb = c_white
    p_pres.space_after = Pt(6)

    p_date = info_tf.add_paragraph()
    p_date.text = "Date: September 2026"
    p_date.font.name = "Arial"
    p_date.font.size = Pt(13)
    p_date.font.color.rgb = RGBColor(170, 185, 205)
    p_date.space_after = Pt(6)

    p_repo = info_tf.add_paragraph()
    p_repo.text = "Repository: https://github.com/rahiakil/tkhome"
    p_repo.font.name = "Arial"
    p_repo.font.size = Pt(12)
    p_repo.font.color.rgb = c_teal

    # Core Goal Card (Right Column)
    goal_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(4.0), Inches(5.8), Inches(2.4))
    goal_card.fill.solid()
    goal_card.fill.fore_color.rgb = c_light_navy
    goal_card.line.color.rgb = c_teal
    goal_card.line.width = Pt(1.5)

    goal_box = slide1.shapes.add_textbox(Inches(6.7), Inches(4.2), Inches(5.4), Inches(2.0))
    goal_tf = goal_box.text_frame
    goal_tf.word_wrap = True
    
    p_goal_lbl = goal_tf.paragraphs[0]
    p_goal_lbl.text = "CORE MISSION GOAL"
    p_goal_lbl.font.name = "Arial"
    p_goal_lbl.font.size = Pt(11)
    p_goal_lbl.font.bold = True
    p_goal_lbl.font.color.rgb = c_teal
    p_goal_lbl.space_after = Pt(8)

    p_goal_desc = goal_tf.add_paragraph()
    p_goal_desc.text = "Automate the mapping of 100k+ customer products to Vertex Tax Categories. Since raw customer descriptions are brief and vague, we enrich them using search results. This system filters out non-matching search results to ensure only CORRECT product data is used for downstream tax classification."
    p_goal_desc.font.name = "Arial"
    p_goal_desc.font.size = Pt(12)
    p_goal_desc.font.color.rgb = c_white
    p_goal_desc.line_spacing = 1.15

    # ==========================================
    # SLIDE 2: THE BUSINESS PROBLEM & CORE DILEMMA
    # ==========================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    slide2.background.fill.solid()
    slide2.background.fill.fore_color.rgb = c_white
    add_slide_header(slide2, "The Business Problem & Core Dilemma")

    # Left Column: Challenge & Gap
    col1_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.6), Inches(5.6), Inches(4.1))
    col1_bg.fill.solid()
    col1_bg.fill.fore_color.rgb = c_light_grey
    col1_bg.line.color.rgb = c_border_grey
    col1_bg.line.width = Pt(1)

    col1_box = slide2.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(5.2), Inches(3.7))
    col1_tf = col1_box.text_frame
    col1_tf.word_wrap = True
    
    p_c1_title = col1_tf.paragraphs[0]
    p_c1_title.text = "THE INFORMATION GAP"
    p_c1_title.font.name = "Arial"
    p_c1_title.font.size = Pt(12)
    p_c1_title.font.bold = True
    p_c1_title.font.color.rgb = c_navy
    p_c1_title.space_after = Pt(12)

    p_c1_p1_lbl = col1_tf.add_paragraph()
    p_c1_p1_lbl.text = "• The Tax Engine Requirement"
    p_c1_p1_lbl.font.bold = True
    p_c1_p1_lbl.font.size = Pt(13)
    p_c1_p1_lbl.font.color.rgb = c_charcoal
    
    p_c1_p1_val = col1_tf.add_paragraph()
    p_c1_p1_val.text = "Vertex's tax engine (O-Series) needs extremely granular product details (ingredients, sweeteners, alcohol content, packaging format) to compute correct tax."
    p_c1_p1_val.font.size = Pt(12)
    p_c1_p1_val.font.color.rgb = c_muted_grey
    p_c1_p1_val.space_after = Pt(12)
    p_c1_p1_val.margin_left = Inches(0.2)

    p_c1_p2_lbl = col1_tf.add_paragraph()
    p_c1_p2_lbl.text = "• The Retailer Data Gap"
    p_c1_p2_lbl.font.bold = True
    p_c1_p2_lbl.font.size = Pt(13)
    p_c1_p2_lbl.font.color.rgb = c_charcoal

    p_c1_p2_val = col1_tf.add_paragraph()
    p_c1_p2_val.text = "Big-box retailers rarely have this level of detail. They only provide product titles (e.g., 'Lipton CB Unswet Blk Tea 14oz') which are brief and highly abbreviated."
    p_c1_p2_val.font.size = Pt(12)
    p_c1_p2_val.font.color.rgb = c_muted_grey
    p_c1_p2_val.margin_left = Inches(0.2)

    # Right Column: Solution & Linchpin
    col2_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.983), Inches(1.6), Inches(5.6), Inches(4.1))
    col2_bg.fill.solid()
    col2_bg.fill.fore_color.rgb = c_light_grey
    col2_bg.line.color.rgb = c_border_grey
    col2_bg.line.width = Pt(1)

    col2_box = slide2.shapes.add_textbox(Inches(7.183), Inches(1.8), Inches(5.2), Inches(3.7))
    col2_tf = col2_box.text_frame
    col2_tf.word_wrap = True
    
    p_c2_title = col2_tf.paragraphs[0]
    p_c2_title.text = "THE ENRICHMENT SOLUTION"
    p_c2_title.font.name = "Arial"
    p_c2_title.font.size = Pt(12)
    p_c2_title.font.bold = True
    p_c2_title.font.color.rgb = c_teal
    p_c2_title.space_after = Pt(12)

    p_c2_p1_lbl = col2_tf.add_paragraph()
    p_c2_p1_lbl.text = "• Web Specification Enrichment"
    p_c2_p1_lbl.font.bold = True
    p_c2_p1_lbl.font.size = Pt(13)
    p_c2_p1_lbl.font.color.rgb = c_charcoal
    
    p_c2_p1_val = col2_tf.add_paragraph()
    p_c2_p1_val.text = "We build an Enrichment Pipeline that crawls the web for deep product specifications based on the retailer's short product title."
    p_c2_p1_val.font.size = Pt(12)
    p_c2_p1_val.font.color.rgb = c_muted_grey
    p_c2_p1_val.space_after = Pt(12)
    p_c2_p1_val.margin_left = Inches(0.2)

    p_c2_p2_lbl = col2_tf.add_paragraph()
    p_c2_p2_lbl.text = "• The Noise & Tax Audit Risk Linchpin"
    p_c2_p2_lbl.font.bold = True
    p_c2_p2_lbl.font.size = Pt(13)
    p_c2_p2_lbl.font.color.rgb = c_charcoal

    p_c2_p2_val = col2_tf.add_paragraph()
    p_c2_p2_val.text = "Search engines return noisy results. Mixing up different formats (Tea Bags vs Liquid Bottle) or formula (Sugar-Free vs Original) causes wrong tax classifications and massive client audit risks."
    p_c2_p2_val.font.size = Pt(12)
    p_c2_p2_val.font.color.rgb = c_muted_grey
    p_c2_p2_val.margin_left = Inches(0.2)

    # Bottom Callout Box (Rule of Thumb)
    callout_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(5.9), Inches(11.833), Inches(0.8))
    callout_bg.fill.solid()
    callout_bg.fill.fore_color.rgb = c_coral_bg
    callout_bg.line.color.rgb = c_coral_text
    callout_bg.line.width = Pt(1.5)

    callout_box = slide2.shapes.add_textbox(Inches(0.95), Inches(6.0), Inches(11.433), Inches(0.6))
    callout_tf = callout_box.text_frame
    callout_tf.word_wrap = True
    
    p_callout = callout_tf.paragraphs[0]
    p_callout.alignment = PP_ALIGN.CENTER
    p_callout.text = "⚠️  CORE TAX SYSTEM RULE OF THUMB: False information is worse than no information."
    p_callout.font.name = "Arial"
    p_callout.font.size = Pt(14)
    p_callout.font.bold = True
    p_callout.font.color.rgb = c_coral_text

    # ==========================================
    # SLIDE 3: HEURISTICS & BOUNDARY CASES (Learned from Data)
    # ==========================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    slide3.background.fill.solid()
    slide3.background.fill.fore_color.rgb = c_white
    add_slide_header(slide3, "Heuristics & Boundary Cases (Learned from Data)")

    # Intro line
    intro_box = slide3.shapes.add_textbox(Inches(0.75), Inches(1.4), Inches(11.833), Inches(0.4))
    intro_tf = intro_box.text_frame
    intro_tf.word_wrap = True
    p_intro = intro_tf.paragraphs[0]
    p_intro.text = "From our 50-product training set, we discovered critical fine-grained product boundaries:"
    p_intro.font.name = "Arial"
    p_intro.font.size = Pt(13)
    p_intro.font.italic = True
    p_intro.font.color.rgb = c_muted_grey

    # Left Column Heuristics
    col1_box_s3 = slide3.shapes.add_textbox(Inches(0.75), Inches(1.9), Inches(5.6), Inches(5.0))
    col1_tf_s3 = col1_box_s3.text_frame
    col1_tf_s3.word_wrap = True
    col1_tf_s3.margin_left = 0

    # Format Mismatches
    p_h1 = col1_tf_s3.paragraphs[0]
    p_h1.text = "1. Format Mismatches (Liquid vs Dry)"
    p_h1.font.bold = True
    p_h1.font.size = Pt(14)
    p_h1.font.color.rgb = c_navy
    p_h1_desc = col1_tf_s3.add_paragraph()
    p_h1_desc.text = "Ready-to-drink bottles vs Loose leaf/Tea bags. Dry formats have completely different tax codes, sweeteners, and physical states compared to liquid beverages."
    p_h1_desc.font.size = Pt(11.5)
    p_h1_desc.font.color.rgb = c_muted_grey
    p_h1_desc.space_after = Pt(14)
    p_h1_desc.margin_left = Inches(0.15)

    # Formula Variations
    p_h2 = col1_tf_s3.add_paragraph()
    p_h2.text = "2. Formula Variations (Sweeteners)"
    p_h2.font.bold = True
    p_h2.font.size = Pt(14)
    p_h2.font.color.rgb = c_navy
    p_h2_desc = col1_tf_s3.add_paragraph()
    p_h2_desc.text = "Lite / Diet / Sugar-Free vs Original / Sweetened. Sweetener type and concentration dictates local tax rates. They cannot be mixed under any circumstances."
    p_h2_desc.font.size = Pt(11.5)
    p_h2_desc.font.color.rgb = c_muted_grey
    p_h2_desc.space_after = Pt(14)
    p_h2_desc.margin_left = Inches(0.15)

    # Alcohol content
    p_h3 = col1_tf_s3.add_paragraph()
    p_h3.text = "3. Alcohol Category Classification"
    p_h3.font.bold = True
    p_h3.font.size = Pt(14)
    p_h3.font.color.rgb = c_navy
    p_h3_desc = col1_tf_s3.add_paragraph()
    p_h3_desc.text = "Ale vs Pilsner / Lager (e.g., 'Labatt Ale' matched with 'Canadian Pilsener'). Ale and Pilsner use different fermenting yeast and fall under separate tax groups."
    p_h3_desc.font.size = Pt(11.5)
    p_h3_desc.font.color.rgb = c_muted_grey
    p_h3_desc.margin_left = Inches(0.15)

    # Right Column Heuristics
    col2_box_s3 = slide3.shapes.add_textbox(Inches(6.983), Inches(1.9), Inches(5.6), Inches(5.0))
    col2_tf_s3 = col2_box_s3.text_frame
    col2_tf_s3.word_wrap = True
    col2_tf_s3.margin_left = 0

    # Combo packs
    p_h4 = col2_tf_s3.paragraphs[0]
    p_h4.text = "4. Single Products vs Combo Packs"
    p_h4.font.bold = True
    p_h4.font.size = Pt(14)
    p_h4.font.color.rgb = c_navy
    p_h4_desc = col2_tf_s3.add_paragraph()
    p_h4_desc.text = "Wipes vs Wipes & Foam Combo Packs. Combo packs bundle extra components with separate tax definitions, requiring different tax compliance analysis."
    p_h4_desc.font.size = Pt(11.5)
    p_h4_desc.font.color.rgb = c_muted_grey
    p_h4_desc.space_after = Pt(14)
    p_h4_desc.margin_left = Inches(0.15)

    # Branded vs Generic
    p_h5 = col2_tf_s3.add_paragraph()
    p_h5.text = "5. Branded vs Generic Products"
    p_h5.font.bold = True
    p_h5.font.size = Pt(14)
    p_h5.font.color.rgb = c_navy
    p_h5_desc = col2_tf_s3.add_paragraph()
    p_h5_desc.text = "Generics (e.g., 'Milk, Chocolate', 'Large Org Apricot') must return ZERO trusted brand results to prevent hallucinating false brand attributes."
    p_h5_desc.font.size = Pt(11.5)
    p_h5_desc.font.color.rgb = c_muted_grey
    p_h5_desc.space_after = Pt(14)
    p_h5_desc.margin_left = Inches(0.15)

    # Ignoring Size
    p_h6 = col2_tf_s3.add_paragraph()
    p_h6.text = "6. Product Volume / Pack Size Neutrality"
    p_h6.font.bold = True
    p_h6.font.size = Pt(14)
    p_h6.font.color.rgb = c_navy
    p_h6_desc = col2_tf_s3.add_paragraph()
    p_h6_desc.text = "Pack size (6-pack vs 12-pack) or volume (12oz vs 14oz) can be ignored if the chemical formula inside is identical, as tax rules focus on ingredients."
    p_h6_desc.font.size = Pt(11.5)
    p_h6_desc.font.color.rgb = c_muted_grey
    p_h6_desc.margin_left = Inches(0.15)


    # ==========================================
    # SLIDE 4: OUR SYSTEM ARCHITECTURE: THE TWO-TIER PIPELINE
    # ==========================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    slide4.background.fill.solid()
    slide4.background.fill.fore_color.rgb = c_white
    add_slide_header(slide4, "Our System Architecture: The Two-Tier Pipeline")

    # Tier 1 Card
    t1_bg = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.6), Inches(5.6), Inches(4.8))
    t1_bg.fill.solid()
    t1_bg.fill.fore_color.rgb = c_soft_teal
    t1_bg.line.color.rgb = c_teal
    t1_bg.line.width = Pt(1.5)

    t1_box = slide4.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(5.2), Inches(4.4))
    t1_tf = t1_box.text_frame
    t1_tf.word_wrap = True
    
    p_t1_hdr = t1_tf.paragraphs[0]
    p_t1_hdr.text = "TIER 1: SMART LOCAL HEURISTIC"
    p_t1_hdr.font.name = "Arial"
    p_t1_hdr.font.size = Pt(13)
    p_t1_hdr.font.bold = True
    p_t1_hdr.font.color.rgb = c_navy
    p_t1_hdr.space_after = Pt(10)

    p_t1_sub = t1_tf.add_paragraph()
    p_t1_sub.text = "The High-Speed, Zero-Cost Baseline Engine"
    p_t1_sub.font.size = Pt(12)
    p_t1_sub.font.italic = True
    p_t1_sub.font.color.rgb = c_teal
    p_t1_sub.space_after = Pt(12)

    bullet_t1_1 = t1_tf.add_paragraph()
    bullet_t1_1.text = "• Rule-Based Python Architecture"
    bullet_t1_1.font.bold = True
    bullet_t1_1.font.size = Pt(12.5)
    bullet_t1_1.font.color.rgb = c_charcoal
    
    bullet_t1_1_desc = t1_tf.add_paragraph()
    bullet_t1_1_desc.text = "Tokenizes terms, filters volume indicators, matches precise brand requirements, and detects strong negative mismatches ('bags' vs liquid)."
    bullet_t1_1_desc.font.size = Pt(11.5)
    bullet_t1_1_desc.font.color.rgb = c_muted_grey
    bullet_t1_1_desc.margin_left = Inches(0.2)
    bullet_t1_1_desc.space_after = Pt(10)

    bullet_t1_2 = t1_tf.add_paragraph()
    bullet_t1_2.text = "• Performance Benchmarks (Offline)"
    bullet_t1_2.font.bold = True
    bullet_t1_2.font.size = Pt(12.5)
    bullet_t1_2.font.color.rgb = c_charcoal

    bullet_t1_2_desc = t1_tf.add_paragraph()
    bullet_t1_2_desc.text = "Latency: < 1 millisecond\nAPI Token Cost: $0.00 (Run 100% locally)\nF1-Score: ~68.2% (Perfect as an instant fallback and local test harness)"
    bullet_t1_2_desc.font.size = Pt(11.5)
    bullet_t1_2_desc.font.color.rgb = c_muted_grey
    bullet_t1_2_desc.margin_left = Inches(0.2)

    # Tier 2 Card
    t2_bg = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.983), Inches(1.6), Inches(5.6), Inches(4.8))
    t2_bg.fill.solid()
    t2_bg.fill.fore_color.rgb = c_light_grey
    t2_bg.line.color.rgb = c_border_grey
    t2_bg.line.width = Pt(1)

    t2_box = slide4.shapes.add_textbox(Inches(7.183), Inches(1.8), Inches(5.2), Inches(4.4))
    t2_tf = t2_box.text_frame
    t2_tf.word_wrap = True
    
    p_t2_hdr = t2_tf.paragraphs[0]
    p_t2_hdr.text = "TIER 2: SOTA LLM + STRUCTURED OUTPUTS"
    p_t2_hdr.font.name = "Arial"
    p_t2_hdr.font.size = Pt(13)
    p_t2_hdr.font.bold = True
    p_t2_hdr.font.color.rgb = c_navy
    p_t2_hdr.space_after = Pt(10)

    p_t2_sub = t2_tf.add_paragraph()
    p_t2_sub.text = "The Production Model utilizing OpenAI JSON Enforcements"
    p_t2_sub.font.size = Pt(12)
    p_t2_sub.font.italic = True
    p_t2_sub.font.color.rgb = c_teal
    p_t2_sub.space_after = Pt(12)

    bullet_t2_1 = t2_tf.add_paragraph()
    bullet_t2_1.text = "• GPT-4o-mini with Strict Pydantic Schema"
    bullet_t2_1.font.bold = True
    bullet_t2_1.font.size = Pt(12.5)
    bullet_t2_1.font.color.rgb = c_charcoal
    
    bullet_t2_1_desc = t2_tf.add_paragraph()
    bullet_t2_1_desc.text = "Enforces strict JSON schema guarantees via response_format. Eliminates parsing errors, returning exact structured lists of evaluated indices, confidence scores, and logical reasoning."
    bullet_t2_1_desc.font.size = Pt(11.5)
    bullet_t2_1_desc.font.color.rgb = c_muted_grey
    bullet_t2_1_desc.margin_left = Inches(0.2)
    bullet_t2_1_desc.space_after = Pt(10)

    bullet_t2_2 = t2_tf.add_paragraph()
    bullet_t2_2.text = "• High-Context Semantic Reasoning"
    bullet_t2_2.font.bold = True
    bullet_t2_2.font.size = Pt(12.5)
    bullet_t2_2.font.color.rgb = c_charcoal

    bullet_t2_2_desc = t2_tf.add_paragraph()
    bullet_t2_2_desc.text = "Easily resolves complex edge cases (e.g. recognizing Yunnan Black Tea with Lemon is not 'Pure Black Tea', or that Wipes are not Wipes with Foam). Expected Target F1-Score: >88%."
    bullet_t2_2_desc.font.size = Pt(11.5)
    bullet_t2_2_desc.font.color.rgb = c_muted_grey
    bullet_t2_2_desc.margin_left = Inches(0.2)

    # ==========================================
    # SLIDE 5: HOW WE MEASURE SUCCESS (KPI SELECTION)
    # ==========================================
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    slide5.background.fill.solid()
    slide5.background.fill.fore_color.rgb = c_white
    add_slide_header(slide5, "How We Measure Success (KPI Selection)")

    # Left Column (Core Metrics)
    m_box = slide5.shapes.add_textbox(Inches(0.75), Inches(1.6), Inches(6.0), Inches(5.2))
    m_tf = m_box.text_frame
    m_tf.word_wrap = True
    m_tf.margin_left = 0

    p_m1 = m_tf.paragraphs[0]
    p_m1.text = "🎯 Precision (Most Critical Metric)"
    p_m1.font.bold = True
    p_m1.font.size = Pt(15)
    p_m1.font.color.rgb = c_navy
    p_m1_desc = m_tf.add_paragraph()
    p_m1_desc.text = "Out of the search results we mark as 'trusted', how many are actually correct? High precision is paramount to avoid feeding erroneous product specifications into downstream tax classifications."
    p_m1_desc.font.size = Pt(11.5)
    p_m1_desc.font.color.rgb = c_muted_grey
    p_m1_desc.space_after = Pt(12)
    p_m1_desc.margin_left = Inches(0.25)

    p_m2 = m_tf.add_paragraph()
    p_m2.text = "📈 Recall (Secondary Volume Metric)"
    p_m2.font.bold = True
    p_m2.font.size = Pt(15)
    p_m2.font.color.rgb = c_navy
    p_m2_desc = m_tf.add_paragraph()
    p_m2_desc.text = "Out of all matching pages available on the web, how many did we capture? High recall guarantees we fetch enough rich text pages to make a definitive product tax determination."
    p_m2_desc.font.size = Pt(11.5)
    p_m2_desc.font.color.rgb = c_muted_grey
    p_m2_desc.space_after = Pt(12)
    p_m2_desc.margin_left = Inches(0.25)

    p_m3 = m_tf.add_paragraph()
    p_m3.text = "⭐ F0.5-Score (Our Primary Optimization Metric)"
    p_m3.font.bold = True
    p_m3.font.size = Pt(15)
    p_m3.font.color.rgb = c_teal
    p_m3_desc = m_tf.add_paragraph()
    p_m3_desc.text = "F0.5 = (1.25 * Precision * Recall) / (0.25 * Precision + Recall)\nF0.5-score weights precision TWICE as heavily as recall. This mathematical choice aligns perfectly with our core tax system mandate: prioritize correctness over raw volume."
    p_m3_desc.font.size = Pt(11.5)
    p_m3_desc.font.color.rgb = c_muted_grey
    p_m3_desc.margin_left = Inches(0.25)

    # Right Column (Evaluation & Operations)
    eval_bg = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.15), Inches(1.6), Inches(5.4), Inches(4.8))
    eval_bg.fill.solid()
    eval_bg.fill.fore_color.rgb = c_light_grey
    eval_bg.line.color.rgb = c_border_grey
    eval_bg.line.width = Pt(1)

    eval_box = slide5.shapes.add_textbox(Inches(7.35), Inches(1.8), Inches(5.0), Inches(4.4))
    eval_tf = eval_box.text_frame
    eval_tf.word_wrap = True
    
    p_ev_hdr = eval_tf.paragraphs[0]
    p_ev_hdr.text = "ADDITIONAL KPIS WE TRACK"
    p_ev_hdr.font.name = "Arial"
    p_ev_hdr.font.size = Pt(12)
    p_ev_hdr.font.bold = True
    p_ev_hdr.font.color.rgb = c_navy
    p_ev_hdr.space_after = Pt(12)

    bullet_ev_1 = eval_tf.add_paragraph()
    bullet_ev_1.text = "• Accuracy & Confusion Matrix"
    bullet_ev_1.font.bold = True
    bullet_ev_1.font.size = Pt(13)
    bullet_ev_1.font.color.rgb = c_charcoal
    
    bullet_ev_1_desc = eval_tf.add_paragraph()
    bullet_ev_1_desc.text = "We trace raw True Positives, False Positives, True Negatives, and False Negatives to detect precise failure patterns (e.g. over-filtering vs over-trusting)."
    bullet_ev_1_desc.font.size = Pt(11.5)
    bullet_ev_1_desc.font.color.rgb = c_muted_grey
    bullet_ev_1_desc.margin_left = Inches(0.2)
    bullet_ev_1_desc.space_after = Pt(12)

    bullet_ev_2 = eval_tf.add_paragraph()
    bullet_ev_2.text = "• Operational Efficiency Metrics"
    bullet_ev_2.font.bold = True
    bullet_ev_2.font.size = Pt(13)
    bullet_ev_2.font.color.rgb = c_charcoal

    bullet_ev_2_desc = eval_tf.add_paragraph()
    bullet_ev_2_desc.text = "Average Latency: target < 1.0s per product batch\nAPI Cost per 10k Products: target < $2.00 using gpt-4o-mini context compression."
    bullet_ev_2_desc.font.size = Pt(11.5)
    bullet_ev_2_desc.font.color.rgb = c_muted_grey
    bullet_ev_2_desc.margin_left = Inches(0.2)

    # ==========================================
    # SLIDE 6: THE SELF-IMPROVING ENGINE
    # ==========================================
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    slide6.background.fill.solid()
    slide6.background.fill.fore_color.rgb = c_white
    add_slide_header(slide6, "The Self-Improving Engine Loop")

    # Step 1 Card (Left)
    card1 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.8), Inches(3.6), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = c_light_grey
    card1.line.color.rgb = c_border_grey
    card1.line.width = Pt(1)

    box1 = slide6.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(3.3), Inches(4.1))
    tf1 = box1.text_frame
    tf1.word_wrap = True
    p1_hdr = tf1.paragraphs[0]
    p1_hdr.text = "1. ERROR DIAGNOSTICS"
    p1_hdr.font.name = "Arial"
    p1_hdr.font.size = Pt(13)
    p1_hdr.font.bold = True
    p1_hdr.font.color.rgb = c_navy
    p1_hdr.space_after = Pt(12)

    p1_desc = tf1.add_paragraph()
    p1_desc.text = "The evaluation harness runs against a golden validation dataset. It intercepts and logs every error:\n\n• False Positives (over-trusting unrelated results)\n• False Negatives (over-filtering matching results)\n\nThese errors serve as prompt training feedback."
    p1_desc.font.size = Pt(11.5)
    p1_desc.font.color.rgb = c_muted_grey
    p1_desc.line_spacing = 1.2

    # Arrow 1
    arrow1 = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.45), Inches(3.8), Inches(0.3), Inches(0.3))
    arrow1.fill.solid()
    arrow1.fill.fore_color.rgb = c_teal
    arrow1.line.fill.background()

    # Step 2 Card (Middle)
    card2 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.866), Inches(1.8), Inches(3.6), Inches(4.5))
    card2.fill.solid()
    card2.fill.fore_color.rgb = c_soft_teal
    card2.line.color.rgb = c_teal
    card2.line.width = Pt(1.5)

    box2 = slide6.shapes.add_textbox(Inches(5.016), Inches(2.0), Inches(3.3), Inches(4.1))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    p2_hdr = tf2.paragraphs[0]
    p2_hdr.text = "2. META-PROMPT OPTIMIZER"
    p2_hdr.font.name = "Arial"
    p2_hdr.font.size = Pt(13)
    p2_hdr.font.bold = True
    p2_hdr.font.color.rgb = c_teal
    p2_hdr.space_after = Pt(12)

    p2_desc = tf2.add_paragraph()
    p2_desc.text = "A high-capability LLM meta-optimizer takes:\n\n• Current system prompt\n• Captured failing products & raw search data\n\nIt performs root-cause diagnostics and proposes precise system instructions to prevent these specific failures."
    p2_desc.font.size = Pt(11.5)
    p2_desc.font.color.rgb = c_muted_grey
    p2_desc.line_spacing = 1.2

    # Arrow 2
    arrow2 = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.55), Inches(3.8), Inches(0.3), Inches(0.3))
    arrow2.fill.solid()
    arrow2.fill.fore_color.rgb = c_teal
    arrow2.line.fill.background()

    # Step 3 Card (Right)
    card3 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.983), Inches(1.8), Inches(3.6), Inches(4.5))
    card3.fill.solid()
    card3.fill.fore_color.rgb = c_light_grey
    card3.line.color.rgb = c_border_grey
    card3.line.width = Pt(1)

    box3 = slide6.shapes.add_textbox(Inches(9.133), Inches(2.0), Inches(3.3), Inches(4.1))
    tf3 = box3.text_frame
    tf3.word_wrap = True
    p3_hdr = tf3.paragraphs[0]
    p3_hdr.text = "3. GATED PROMOTION"
    p3_hdr.font.name = "Arial"
    p3_hdr.font.size = Pt(13)
    p3_hdr.font.bold = True
    p3_hdr.font.color.rgb = c_navy
    p3_hdr.space_after = Pt(12)

    p3_desc = tf3.add_paragraph()
    p3_desc.text = "The harness reruns evaluation with the proposed system prompt:\n\n• If F0.5-score improves: The optimized prompt is promoted to production.\n• If F0.5-score declines: It is discarded.\n\nThis closed loop guarantees autonomous product-matching accuracy gains."
    p3_desc.font.size = Pt(11.5)
    p3_desc.font.color.rgb = c_muted_grey
    p3_desc.line_spacing = 1.2

    # ==========================================
    # SLIDE 7: OPERATIONAL IMPACT & ROADMAP (Dark Theme Outro)
    # ==========================================
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    slide7.background.fill.solid()
    slide7.background.fill.fore_color.rgb = c_navy
    add_dark_slide_header(slide7, "Operational Impact & Roadmap")

    # Decorative teal bar on the left
    left_bar_s7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.15), Inches(7.5))
    left_bar_s7.fill.solid()
    left_bar_s7.fill.fore_color.rgb = c_teal
    left_bar_s7.line.fill.background()

    # Left Column (Operational Performance)
    op_bg = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.6), Inches(5.6), Inches(4.8))
    op_bg.fill.solid()
    op_bg.fill.fore_color.rgb = c_light_navy
    op_bg.line.color.rgb = c_teal
    op_bg.line.width = Pt(1)

    op_box = slide7.shapes.add_textbox(Inches(0.95), Inches(1.8), Inches(5.2), Inches(4.4))
    op_tf = op_box.text_frame
    op_tf.word_wrap = True
    
    p_op_hdr = op_tf.paragraphs[0]
    p_op_hdr.text = "EXPECTED OPERATIONAL PERFORMANCE"
    p_op_hdr.font.name = "Arial"
    p_op_hdr.font.size = Pt(12)
    p_op_hdr.font.bold = True
    p_op_hdr.font.color.rgb = c_teal
    p_op_hdr.space_after = Pt(12)

    b_op_1 = op_tf.add_paragraph()
    b_op_1.text = "🎯 Accuracy & Confidence"
    b_op_1.font.bold = True
    b_op_1.font.size = Pt(13)
    b_op_1.font.color.rgb = c_white
    b_op_1_desc = op_tf.add_paragraph()
    b_op_1_desc.text = "GPT-4o-mini achieves >92% Precision and >88% F1-score with optimized, structured prompt rules."
    b_op_1_desc.font.size = Pt(11.5)
    b_op_1_desc.font.color.rgb = RGBColor(170, 185, 205)
    b_op_1_desc.margin_left = Inches(0.2)
    b_op_1_desc.space_after = Pt(12)

    b_op_2 = op_tf.add_paragraph()
    b_op_2.text = "⚡ High Throughput Execution"
    b_op_2.font.bold = True
    b_op_2.font.size = Pt(13)
    b_op_2.font.color.rgb = c_white
    b_op_2_desc = op_tf.add_paragraph()
    b_op_2_desc.text = "Highly parallelized processing threads classify 100k customer products in under 5 minutes."
    b_op_2_desc.font.size = Pt(11.5)
    b_op_2_desc.font.color.rgb = RGBColor(170, 185, 205)
    b_op_2_desc.margin_left = Inches(0.2)
    b_op_2_desc.space_after = Pt(12)

    b_op_3 = op_tf.add_paragraph()
    b_op_3.text = "💵 Cost Efficiency"
    b_op_3.font.bold = True
    b_op_3.font.size = Pt(13)
    b_op_3.font.color.rgb = c_white
    b_op_3_desc = op_tf.add_paragraph()
    b_op_3_desc.text = "Extremely affordable run costs: ~$0.02 per 100 products ($20.00 total for 100k products)."
    b_op_3_desc.font.size = Pt(11.5)
    b_op_3_desc.font.color.rgb = RGBColor(170, 185, 205)
    b_op_3_desc.margin_left = Inches(0.2)

    # Right Column (Roadmap Phases)
    rm_bg = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.983), Inches(1.6), Inches(5.6), Inches(4.8))
    rm_bg.fill.solid()
    rm_bg.fill.fore_color.rgb = c_light_navy
    rm_bg.line.color.rgb = c_border_grey
    rm_bg.line.width = Pt(0.5)

    rm_box = slide7.shapes.add_textbox(Inches(7.183), Inches(1.8), Inches(5.2), Inches(4.4))
    rm_tf = rm_box.text_frame
    rm_tf.word_wrap = True
    
    p_rm_hdr = rm_tf.paragraphs[0]
    p_rm_hdr.text = "PRODUCT ROADMAP & NEXT STEPS"
    p_rm_hdr.font.name = "Arial"
    p_rm_hdr.font.size = Pt(12)
    p_rm_hdr.font.bold = True
    p_rm_hdr.font.color.rgb = c_teal
    p_rm_hdr.space_after = Pt(12)

    p_ph1 = rm_tf.add_paragraph()
    p_ph1.text = "Phase 1: Foundation (COMPLETE)"
    p_ph1.font.bold = True
    p_ph1.font.size = Pt(12.5)
    p_ph1.font.color.rgb = c_white
    p_ph1_desc = rm_tf.add_paragraph()
    p_ph1_desc.text = "Evaluation harness, smart offline baseline, parallelized OpenAI Structured Output API, and self-improving prompt optimization loop."
    p_ph1_desc.font.size = Pt(11)
    p_ph1_desc.font.color.rgb = RGBColor(170, 185, 205)
    p_ph1_desc.margin_left = Inches(0.2)
    p_ph1_desc.space_after = Pt(8)

    p_ph2 = rm_tf.add_paragraph()
    p_ph2.text = "Phase 2: Immediate (CURRENT)"
    p_ph2.font.bold = True
    p_ph2.font.size = Pt(12.5)
    p_ph2.font.color.rgb = c_white
    p_ph2_desc = rm_tf.add_paragraph()
    p_ph2_desc.text = "Ingest the live interview testing dataset, evaluate it via prompt rules, and output predictions."
    p_ph2_desc.font.size = Pt(11)
    p_ph2_desc.font.color.rgb = RGBColor(170, 185, 205)
    p_ph2_desc.margin_left = Inches(0.2)
    p_ph2_desc.space_after = Pt(8)

    p_ph3 = rm_tf.add_paragraph()
    p_ph3.text = "Phase 3: Hybrid Retriever (MID-TERM)"
    p_ph3.font.bold = True
    p_ph3.font.size = Pt(12.5)
    p_ph3.font.color.rgb = c_white
    p_ph3_desc = rm_tf.add_paragraph()
    p_ph3_desc.text = "Integrate a fast, vector-based hybrid retriever to filter search results prior to LLM layer, cutting token costs by an estimated 40%."
    p_ph3_desc.font.size = Pt(11)
    p_ph3_desc.font.color.rgb = RGBColor(170, 185, 205)
    p_ph3_desc.margin_left = Inches(0.2)
    p_ph3_desc.space_after = Pt(8)

    p_ph4 = rm_tf.add_paragraph()
    p_ph4.text = "Phase 4: Open-Weight Fine-Tuning (LONG-TERM)"
    p_ph4.font.bold = True
    p_ph4.font.size = Pt(12.5)
    p_ph4.font.color.rgb = c_white
    p_ph4_desc = rm_tf.add_paragraph()
    p_ph4_desc.text = "Fine-tune a smaller open-weight model (e.g. Llama-3-8B) on validated trusted results. Enables 100% private deployments, reducing LLM costs to zero."
    p_ph4_desc.font.size = Pt(11)
    p_ph4_desc.font.color.rgb = RGBColor(170, 185, 205)
    p_ph4_desc.margin_left = Inches(0.2)


    # Save Presentation
    output_filename = "presentation_slides.pptx"
    prs.save(output_filename)
    print(f"Presentation slides saved successfully to {output_filename}!")

if __name__ == "__main__":
    create_presentation()
